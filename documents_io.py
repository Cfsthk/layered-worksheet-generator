"""Bounded local extraction and native, editable Word output."""
import base64
import io
import re
import shutil
import subprocess
import sys
import tempfile
import zipfile
from pathlib import Path
from PIL import Image, ImageDraw, ImageOps
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Mm, Pt, RGBColor
from domain import AppError, worksheet, question

MAX_FILE = 15 * 1024 * 1024
Image.MAX_IMAGE_PIXELS = 24_000_000
SUBJECTS = {"maths": "數學", "chinese": "中文", "english": "英文"}


def image_data(raw):
    try:
        with Image.open(io.BytesIO(raw)) as source:
            if source.width * source.height > 24_000_000:
                raise AppError("圖片太大，請縮小至 2400 萬像素以下。")
            image = ImageOps.exif_transpose(source).convert("RGB")
            image.thumbnail((1800, 1800))
            out = io.BytesIO()
            image.save(out, "JPEG", quality=86)
            return "data:image/jpeg;base64," + base64.b64encode(out.getvalue()).decode()
    except AppError:
        raise
    except Exception:
        raise AppError("未能讀取圖片，請改用 JPG 或 PNG。", "invalid_file") from None


def _xml_text(node):
    name = node.tag.rsplit("}", 1)[-1]
    if name in {"t", "delText"}:
        return node.text or ""
    if name == "tab":
        return "\t"
    if name in {"br", "cr"}:
        return "\n"
    if name == "f":
        num, den = node.find(qn("m:num")), node.find(qn("m:den"))
        if num is not None and den is not None:
            return "(" + _xml_text(num) + ")/(" + _xml_text(den) + ")"
    return "".join(_xml_text(child) for child in node)


def parse_document(name, encoded):
    if not isinstance(name, str) or len(name) > 240 or not isinstance(encoded, str):
        raise AppError("檔案資料格式錯誤。")
    extension = Path(name).suffix.lower()
    if extension not in {".pdf", ".docx", ".png", ".jpg", ".jpeg", ".heic"}:
        raise AppError("支援 PDF、DOCX、JPG、PNG 或 HEIC。", "invalid_file")
    try:
        raw = base64.b64decode(encoded, validate=True)
    except (ValueError, TypeError):
        raise AppError("檔案傳送不完整，請重新選擇。", "invalid_file") from None
    if not raw or len(raw) > MAX_FILE:
        raise AppError("請選擇不超過 15 MB 的檔案。", "file_too_large", 413)
    result = {"fileName": name, "text": "", "images": [], "pages": 1, "notices": []}
    if extension == ".docx":
        try:
            with zipfile.ZipFile(io.BytesIO(raw)) as archive:
                infos = archive.infolist()
                if len(infos) > 2000 or sum(i.file_size for i in infos) > 80 * 1024 * 1024:
                    raise AppError("Word 檔案解壓後太大，請分拆檔案。", "file_too_large", 413)
                if any("vbaproject" in i.filename.lower() for i in infos):
                    raise AppError("不支援含巨集的文件。", "invalid_file")
                if "word/document.xml" not in archive.namelist():
                    raise AppError("這不是有效的 DOCX 文件。", "invalid_file")
            document = Document(io.BytesIO(raw))
            lines = []
            for element in document.element.body:
                if element.tag == qn("w:tbl"):
                    for row in element.findall(qn("w:tr")):
                        lines.append(" | ".join(_xml_text(cell) for cell in row.findall(qn("w:tc"))))
                else:
                    lines.append(_xml_text(element))
            result["text"] = "\n".join(lines).strip()
            for relation in document.part.rels.values():
                if relation.is_external:
                    continue  # Never fetch links, linked images or remote templates.
                if "image" in relation.reltype:
                    if len(result["images"]) >= 8:
                        raise AppError("文件有超過 8 張內嵌圖片，請分拆檔案。", "file_too_large", 413)
                    try:
                        result["images"].append(image_data(relation.target_part.blob))
                    except AppError:
                        result["notices"].append("部分 Word 向量圖像未能讀取。請將文件另存 PDF 後上載，以免遺漏圖表。")
            result["notices"].append("Word 已擷取文字、表格及可讀圖片；版面、浮動物件及自動題號可能不同，請檢查讀取內容。")
        except AppError:
            raise
        except Exception:
            raise AppError("Word 文件損壞或受保護。請另存為 DOCX 或 PDF 後重試。", "invalid_file") from None
    elif extension == ".pdf":
        if not raw.startswith(b"%PDF-"):
            raise AppError("這不是有效的 PDF。", "invalid_file")
        for binary in ["pdfinfo", "pdftoppm"]:
            if not shutil.which(binary):
                raise AppError("本機缺少 Poppler。請按照 README 安裝，或改用 DOCX／圖片。", "dependency_missing", 503)
        try:
            with tempfile.TemporaryDirectory(prefix="seven-pdf-") as folder:
                folder = Path(folder)
                source = folder / "upload.pdf"
                source.write_bytes(raw)
                info = subprocess.run([shutil.which("pdfinfo"), str(source)], capture_output=True, timeout=20)
                if info.returncode:
                    raise AppError("PDF 受密碼保護或無法讀取。請另存沒有密碼的版本。", "invalid_file")
                match = re.search(r"Pages:\s+(\d+)", info.stdout.decode(errors="replace"))
                pages = int(match.group(1)) if match else 0
                if not 1 <= pages <= 8:
                    raise AppError("每次上載支援 1 至 8 頁。請先分拆較長的 PDF。", "file_too_large", 413)
                result["pages"] = pages
                if shutil.which("pdftotext"):
                    extracted = subprocess.run([shutil.which("pdftotext"), "-layout", str(source), str(folder / "text.txt")], capture_output=True, timeout=25)
                    if extracted.returncode == 0:
                        result["text"] = (folder / "text.txt").read_text(errors="replace").strip()
                else:
                    # The desktop bundle has pdfinfo/pdftoppm but may not include
                    # pdftotext. Isolate pypdf extraction in a timed child process.
                    script = "from pypdf import PdfReader; import sys; r=PdfReader(sys.argv[1]); parts=[]\nfor p in r.pages:\n parts.append((p.extract_text() or '')[:60001])\n if sum(map(len,parts))>60000: break\nprint('\\n'.join(parts))"
                    extracted = subprocess.run([sys.executable, "-c", script, str(source)], capture_output=True, timeout=25)
                    if extracted.returncode == 0:
                        result["text"] = extracted.stdout.decode(errors="replace").strip()
                rendered = subprocess.run([shutil.which("pdftoppm"), "-jpeg", "-scale-to", "1800", str(source), str(folder / "page")], capture_output=True, timeout=45)
                files = sorted(folder.glob("page-*.jpg"))
                if rendered.returncode or len(files) != pages:
                    raise AppError("未能讀取所有 PDF 頁面，請改用圖片或較小的檔案。", "invalid_file")
                result["images"] = [image_data(p.read_bytes()) for p in files]
        except subprocess.TimeoutExpired:
            raise AppError("PDF 處理逾時，請分拆或壓縮檔案。", "invalid_file") from None
    elif extension == ".heic":
        if not shutil.which("sips"):
            raise AppError("這部電腦未能轉換 HEIC，請改用 JPG 或 PNG。", "dependency_missing", 503)
        try:
            with tempfile.TemporaryDirectory(prefix="seven-image-") as folder:
                source, target = Path(folder) / "upload.heic", Path(folder) / "image.jpg"
                source.write_bytes(raw)
                converted = subprocess.run([shutil.which("sips"), "-s", "format", "jpeg", str(source), "--out", str(target)], capture_output=True, timeout=20)
                if converted.returncode or not target.exists():
                    raise AppError("未能讀取 HEIC，請改用 JPG 或 PNG。", "invalid_file")
                result["images"] = [image_data(target.read_bytes())]
        except subprocess.TimeoutExpired:
            raise AppError("圖片轉換逾時，請改用 JPG。", "invalid_file") from None
    else:
        result["images"] = [image_data(raw)]
    if len(result["text"]) > 60000:
        raise AppError("文件超過 60,000 字，請分拆後重試。", "file_too_large", 413)
    if not result["text"] and not result["images"]:
        raise AppError("檔案內找不到可讀內容。", "invalid_file")
    return result


def add_math_text(paragraph, content):
    """Editable OMML fractions inside ordinary editable prose."""
    offset = 0
    for match in re.finditer(r"(?<![\w/])(\d+)\s*[/／]\s*(\d+)(?![\w/])", content):
        paragraph.add_run(content[offset:match.start()])
        math = OxmlElement("m:oMath")
        fraction = OxmlElement("m:f")
        for tag, value in zip(["m:num", "m:den"], match.groups()):
            element = OxmlElement(tag)
            run, node = OxmlElement("m:r"), OxmlElement("m:t")
            node.text = value
            run.append(node)
            element.append(run)
            fraction.append(element)
        math.append(fraction)
        paragraph._p.append(math)
        offset = match.end()
    paragraph.add_run(content[offset:])


def bar_picture(value):
    pairs = value["fractions"]
    image = Image.new("RGB", (1000, len(pairs) * 75), "white")
    draw = ImageDraw.Draw(image)
    for i, (num, den) in enumerate(pairs):
        y = i * 75 + 15
        draw.rectangle((0, y, num / den * 999, y + 42), fill="#aebbe0")
        draw.rectangle((0, y, 999, y + 42), outline="#505050", width=2)
        for tick in range(1, den):
            x = round(tick / den * 999)
            draw.line((x, y, x, y + 42), fill="#505050", width=2)
    out = io.BytesIO()
    image.save(out, "PNG")
    out.seek(0)
    return out


def export_docx(payload):
    if not isinstance(payload, dict) or payload.get("approved") is not True:
        raise AppError("請先檢查這個版本的題目及答案。", "review_required", 409)
    data = worksheet(payload.get("worksheet"), require_answers=False)
    mode = payload.get("mode", "student")
    if mode not in {"student", "answers", "both"}:
        raise AppError("匯出模式錯誤。")
    level = payload.get("level", 4)
    if type(level) is not int or not 1 <= level <= 7:
        raise AppError("程度必須介乎 1 至 7。")
    extra = payload.get("extension")
    if extra:
        extra = {"objective": str(extra.get("objective", ""))[:1000], "question": question(extra.get("question"), 60)}
    doc = Document()
    section = doc.sections[0]
    # Match the existing app's A4 classroom print template.
    section.page_width, section.page_height = Mm(210), Mm(297)
    section.top_margin = section.bottom_margin = Mm(18)
    section.left_margin = section.right_margin = Mm(20)
    for style_name, size in [("Normal", 12), ("Title", 23), ("Heading 1", 14), ("Heading 2", 12)]:
        style = doc.styles[style_name]
        style.font.name, style.font.size = "Arial", Pt(size)
        style.font.color.rgb = RGBColor(0, 0, 0)
        fonts = style.element.get_or_add_rPr().get_or_add_rFonts()
        # Explicit fonts override the stock template's theme fonts. Heiti TC is
        # available on this Mac and renders reliably in Word and headless LO.
        for attr in list(fonts.attrib):
            if "theme" in attr.lower():
                del fonts.attrib[attr]
        fonts.set(qn("w:eastAsia"), "Heiti TC")
        fonts.set(qn("w:cs"), "Arial")
        for color in style.element.xpath("./w:rPr/w:color"):
            for attr in list(color.attrib):
                if "theme" in attr.lower():
                    del color.attrib[attr]
        for border in style.element.xpath("./w:pPr/w:pBdr"):
            border.getparent().remove(border)
        style.paragraph_format.line_spacing = 1.2
        style.paragraph_format.space_after = Pt(6)
    doc.styles["Title"].paragraph_format.space_before = Pt(0)
    doc.styles["Heading 1"].paragraph_format.space_before = Pt(10)
    answer_style = doc.styles.add_style("Answer text", 1)
    answer_style.base_style = doc.styles["Normal"]
    answer_style.font.size = Pt(11)
    answer_style.paragraph_format.line_spacing = 1.1
    answer_style.paragraph_format.space_after = Pt(4)
    doc.core_properties.title = data["topic"]
    doc.core_properties.author = ""
    doc.core_properties.subject = data["objective"]
    footer = section.footer.paragraphs[0]
    footer.alignment = 2
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)

    def add_question(q, number, answers=False):
        prompt = doc.add_paragraph(style="Answer text" if answers else "Normal")
        prompt.paragraph_format.space_before = Pt(6 if answers else 7)
        prompt.paragraph_format.keep_with_next = bool(answers or q["hint"] or q["diagram"] or q["workLines"])
        add_math_text(prompt, f"{number}. {q['prompt']}")
        if answers:
            answer = doc.add_paragraph(style="Answer text")
            answer.paragraph_format.keep_with_next = bool(q["explanation"])
            add_math_text(answer, "答案：" + (q["answer"] or "請老師補充"))
            if q["explanation"]:
                detail = doc.add_paragraph(style="Answer text")
                add_math_text(detail, "解說：" + q["explanation"])
            return
        if q["hint"]:
            cue = doc.add_paragraph()
            cue.paragraph_format.keep_with_next = True
            add_math_text(cue, "提示：" + q["hint"])
        if q["diagram"]:
            graphic = q["diagram"]
            label = doc.add_paragraph()
            label.paragraph_format.keep_with_next = True
            add_math_text(label, "圖解由上至下：" + "、".join(f"{a}/{b}" for a, b in graphic["fractions"]))
            image_p = doc.add_paragraph()
            run = image_p.add_run()
            shape = run.add_picture(bar_picture(graphic), width=Inches(4.8))
            shape._inline.docPr.set("descr", graphic["caption"] or "等長分數圖")
            image_p.paragraph_format.keep_with_next = q["workLines"] > 0
        for line_index in range(q["workLines"]):
            line = doc.add_paragraph(" ")
            line.paragraph_format.line_spacing = Pt(16)
            line.paragraph_format.space_after = Pt(5)
            # Never strand the last answer line on a page of its own.
            line.paragraph_format.keep_with_next = line_index < q["workLines"] - 1
            borders, bottom = OxmlElement("w:pBdr"), OxmlElement("w:bottom")
            for key, value in {"val": "single", "sz": "3", "color": "CCCCCC"}.items():
                bottom.set(qn("w:" + key), value)
            borders.append(bottom)
            between = OxmlElement("w:between")
            for key, value in {"val": "single", "sz": "3", "color": "CCCCCC"}.items():
                between.set(qn("w:" + key), value)
            borders.append(between)
            line._p.get_or_add_pPr().append(borders)

    def add_sheet(answers):
        first = doc.add_paragraph(f"小學{data['grade']}年級 {SUBJECTS[data['subject']]}科")
        if answers and mode == "both":
            # A break-only paragraph can spill onto an empty page when the
            # student sheet fills its last page. Break on the next heading.
            first.paragraph_format.page_break_before = True
        title = data["topic"].strip()
        doc.add_paragraph(title + (" 答案" if answers else ""), "Title")
        doc.add_paragraph(f"程度 {level}    學習目標：{data['objective']}")
        if not answers:
            doc.add_paragraph("姓名：________________    班別：________    日期：____________")
        if data["context"]:
            doc.add_paragraph("閱讀材料", "Heading 1")
            for line in data["context"].splitlines():
                p = doc.add_paragraph()
                add_math_text(p, line)
        previous = None
        for i, q in enumerate(data["questions"]):
            if q["section"] != previous:
                # Comparison symbols and punctuation are part of the teaching
                # instruction, not decoration. Domain validation handles XML.
                doc.add_paragraph(q["section"], "Heading 1")
                previous = q["section"]
            add_question(q, str(i + 1), answers)
        if extra:
            doc.add_paragraph("延伸挑戰", "Heading 1")
            doc.add_paragraph("延伸目標：" + extra["objective"])
            add_question(extra["question"], "E1", answers)
    if mode in {"student", "both"}:
        add_sheet(False)
    if mode in {"answers", "both"}:
        add_sheet(True)
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()
